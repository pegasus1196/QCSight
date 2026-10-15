"""
FastAPI routes for cabinet inspection API.
Handles file uploads, runs CV pipeline, returns results.
"""
import io
import cv2
import numpy as np
from pathlib import Path
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from typing import Optional

from ..cv.inspector import CabinetInspector
from ..models.schemas import InspectionResponse, HealthResponse

router = APIRouter()

UPLOAD_DIR = Path(__file__).parent.parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)


def _read_upload(file: UploadFile) -> np.ndarray:
    contents = file.file.read()
    nparr = np.frombuffer(contents, np.uint8)
    img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=400, detail=f"Cannot decode image: {file.filename}")
    return img


def _parse_label_list(file: UploadFile) -> list[dict]:
    from openpyxl import load_workbook
    contents = file.file.read()
    wb = load_workbook(io.BytesIO(contents), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    results = []
    for row in rows:
        if not row or not row[0]:
            continue
        device = str(row[0]).strip()
        label = str(row[1]).strip() if len(row) > 1 and row[1] else device
        position = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        results.append({"device": device, "label": label, "name": f"{device} ({label})", "id": label})
    return results


def _parse_wire_list(file: UploadFile) -> tuple[list[dict], list[dict]]:
    from openpyxl import load_workbook
    contents = file.file.read()
    wb = load_workbook(io.BytesIO(contents), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    ferrules = []
    colors = []
    for row in rows:
        if not row or not row[0]:
            continue
        wire_id = str(row[0]).strip()
        color = str(row[3]).strip().lower() if len(row) > 3 and row[3] else ""
        ferrule_id = str(row[4]).strip() if len(row) > 4 and row[4] else ""
        if ferrule_id:
            ferrules.append({"id": ferrule_id})
        if color and color in ("red", "blue", "green", "yellow", "black", "white", "brown", "orange", "grey"):
            colors.append({"name": f"Wire {wire_id}", "color": color})
    return ferrules, colors


@router.get("/health", response_model=HealthResponse)
async def health():
    return HealthResponse()


@router.post("/inspect", response_model=InspectionResponse)
async def run_inspection(
    photo: UploadFile = File(..., description="Cabinet photo (JPG/PNG)"),
    reference: Optional[UploadFile] = File(None, description="GA drawing / reference image"),
    wire_list: Optional[UploadFile] = File(None, description="Wire list Excel"),
    label_list: Optional[UploadFile] = File(None, description="Label list Excel"),
    project_name: str = Form("Control cabinet QC"),
    checks: str = Form("presence,labels,ferrules,alignment,colour"),
    confidence_threshold: float = Form(0.72),
):
    photo_img = _read_upload(photo)

    ref_img = None
    if reference:
        try:
            ref_img = _read_upload(reference)
        except HTTPException:
            ref_img = None

    check_list = [c.strip() for c in checks.split(",") if c.strip()]

    expected_labels = None
    expected_devices = None
    if label_list:
        parsed = _parse_label_list(label_list)
        if parsed:
            expected_labels = [{"device": p["device"], "label": p["label"]} for p in parsed]
            expected_devices = [{"name": p["name"], "id": p["id"]} for p in parsed]

    expected_ferrules = None
    expected_colors = None
    if wire_list:
        ferrules, colors = _parse_wire_list(wire_list)
        if ferrules:
            expected_ferrules = ferrules
        if colors:
            expected_colors = colors

    inspector = CabinetInspector(confidence_threshold=confidence_threshold)

    result = inspector.inspect(
        photo=photo_img,
        reference=ref_img,
        checks=check_list,
        project_name=project_name,
        photo_filename=photo.filename or "cabinet.jpg",
        expected_devices=expected_devices,
        expected_labels=expected_labels,
        expected_ferrules=expected_ferrules,
        expected_wire_colors=expected_colors,
    )

    return result


@router.post("/inspect/demo", response_model=InspectionResponse)
async def run_demo_inspection(
    confidence_threshold: float = Form(0.72),
    checks: str = Form("presence,labels,ferrules,alignment,colour"),
):
    demo_path = Path(__file__).parent.parent.parent / "assets" / "cabinet-demo.png"
    ref_path = Path(__file__).parent.parent.parent / "assets" / "cabinet-reference.png"

    if not demo_path.exists():
        raise HTTPException(status_code=404, detail="Demo cabinet image not found")

    photo_img = cv2.imread(str(demo_path))
    if photo_img is None:
        raise HTTPException(status_code=500, detail="Cannot read demo image")

    ref_img = None
    if ref_path.exists():
        ref_img = cv2.imread(str(ref_path))

    check_list = [c.strip() for c in checks.split(",") if c.strip()]

    inspector = CabinetInspector(confidence_threshold=confidence_threshold)

    result = inspector.inspect(
        photo=photo_img,
        reference=ref_img,
        checks=check_list,
        project_name="Control cabinet QC (Demo)",
        photo_filename="cabinet-demo.png",
    )

    return result


@router.post("/analyze/features")
async def analyze_features(
    photo: UploadFile = File(..., description="Cabinet photo to analyze"),
):
    """Returns detected features/keypoints for visualization on frontend."""
    from ..cv.feature_extractor import FeatureExtractor, DeviceDetector

    photo_img = _read_upload(photo)
    h, w = photo_img.shape[:2]

    extractor = FeatureExtractor(method="orb", n_features=1500)
    kps, descs = extractor.extract(photo_img)

    detector = DeviceDetector()
    regions = detector.find_device_contours(photo_img)

    keypoints = [
        {"x": round(kp.pt[0] / w * 100, 2), "y": round(kp.pt[1] / h * 100, 2), "size": round(kp.size, 1)}
        for kp in kps[:200]
    ]

    detected_regions = [
        {
            "x": round(r.x / w * 100, 1), "y": round(r.y / h * 100, 1),
            "w": round(r.w / w * 100, 1), "h": round(r.h / h * 100, 1),
            "confidence": round(r.confidence, 3),
            "keypoints": r.keypoints_count,
        }
        for r in regions
    ]

    return {
        "total_keypoints": len(kps),
        "total_regions": len(regions),
        "descriptor_shape": list(descs.shape) if descs is not None else None,
        "keypoints": keypoints,
        "regions": detected_regions,
    }


@router.post("/analyze/ocr")
async def analyze_ocr(
    photo: UploadFile = File(..., description="Cabinet photo for OCR"),
):
    """Returns all detected text in the image for frontend overlay."""
    from ..cv.ocr_reader import CabinetOCR

    photo_img = _read_upload(photo)
    h, w = photo_img.shape[:2]
    ocr = CabinetOCR()
    results = ocr.read_text(photo_img)

    return {
        "total_detections": len(results),
        "texts": [
            {
                "text": r.text,
                "confidence": round(r.confidence, 3),
                "bbox": {
                    "x": round(r.bbox[0] / w * 100, 1),
                    "y": round(r.bbox[1] / h * 100, 1),
                    "w": round(r.bbox[2] / w * 100, 1),
                    "h": round(r.bbox[3] / h * 100, 1),
                },
                "method": r.method,
            }
            for r in results
        ],
    }


@router.post("/analyze/colors")
async def analyze_colors(
    photo: UploadFile = File(..., description="Cabinet photo for color analysis"),
):
    """Returns dominant color clusters for frontend visualization."""
    from ..cv.color_analyzer import WireColorAnalyzer

    photo_img = _read_upload(photo)
    analyzer = WireColorAnalyzer()

    wire_regions = analyzer.detect_wire_regions(photo_img)
    h, w = photo_img.shape[:2]

    overall = analyzer.detect_dominant_colors(photo_img, n_colors=5)

    region_colors = []
    for rx, ry, rw, rh in wire_regions[:10]:
        roi = photo_img[ry:ry+rh, rx:rx+rw]
        colors = analyzer.detect_dominant_colors(roi, n_colors=2)
        region_colors.append({
            "region": {
                "x": round(rx / w * 100, 1), "y": round(ry / h * 100, 1),
                "w": round(rw / w * 100, 1), "h": round(rh / h * 100, 1),
            },
            "colors": [
                {"name": c.color_name, "confidence": round(c.confidence, 3), "ratio": round(c.pixel_ratio, 3)}
                for c in colors
            ],
        })

    return {
        "dominant_colors": [
            {"name": c.color_name, "confidence": round(c.confidence, 3), "ratio": round(c.pixel_ratio, 3),
             "hsv": list(c.dominant_hsv)}
            for c in overall
        ],
        "wire_regions": region_colors,
    }
