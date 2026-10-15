"""
Generate test input files for CabinetSight real CV pipeline testing.
Creates: GA drawing PDF, wire list Excel, label list Excel.
Use cabinet-demo.png (already in assets/) as the photo input.

These files are crafted so the CV pipeline produces a realistic mix:
  - PASS: labels/devices that exist in the demo image
  - FAIL: labels/devices that don't exist
  - UNREADABLE: borderline detections
"""
import sys
from pathlib import Path

ROOT = Path(__file__).parent.parent
ASSETS = ROOT / "assets" / "test-inputs"
ASSETS.mkdir(parents=True, exist_ok=True)


def create_label_list():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    ws.append(["Device", "Expected Label", "Position"])
    ws.append(["Power Supply", "PS11", "Top-left"])
    ws.append(["Baseplate", "1CBP03", "Center"])
    ws.append(["Drive Unit", "D12", "Top-right"])
    ws.append(["Contactor", "K15", "Mid-left"])
    ws.append(["Terminal Block", "X3", "Bottom-right"])
    ws.append(["PLC Module", "PLC-01", "Mid-center"])
    ws.append(["Safety Relay", "KS2", "Bottom-left"])
    ws.append(["Fuse Holder", "F7", "Top-center"])

    out = ASSETS / "test_label_list.xlsx"
    wb.save(str(out))
    print(f"  Label list: {out} ({ws.max_row - 1} labels)")
    return out


def create_wire_list():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Wires"
    ws.append(["Wire ID", "From", "To", "Color", "Ferrule ID"])
    ws.append(["W001", "PS11:+24V", "CB41:1", "Red", "X3-1"])
    ws.append(["W002", "PS11:0V", "N-bar", "Blue", "X3-2"])
    ws.append(["W003", "CB41:2", "K15:A1", "Red", "K15-A1"])
    ws.append(["W004", "CB42:2", "D12:L1", "Black", "D12-L1"])
    ws.append(["W005", "K15:13", "PLC-01:I0.0", "White", "PLC-I0"])
    ws.append(["W006", "PE-bar", "Cabinet door", "Green/Yellow", "PE-D"])
    ws.append(["W007", "CB43:2", "X3:3", "Brown", "X3-3"])
    ws.append(["W008", "X3:5", "KS2:S11", "Orange", "KS2-S11"])

    out = ASSETS / "test_wire_list.xlsx"
    wb.save(str(out))
    print(f"  Wire list:  {out} ({ws.max_row - 1} wires)")
    return out


def create_ga_drawing():
    from reportlab.lib.pagesizes import A3
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor

    out = ASSETS / "test_ga_drawing.pdf"
    c = canvas.Canvas(str(out), pagesize=A3)
    w, h = A3

    c.setFillColor(HexColor("#f5f5f5"))
    c.rect(0, 0, w, h, fill=1)

    c.setFont("Helvetica-Bold", 18)
    c.setFillColor(HexColor("#1a1a1a"))
    c.drawString(30 * mm, h - 25 * mm, "GENERAL ARRANGEMENT - CONTROL CABINET")

    c.setFont("Helvetica", 10)
    c.setFillColor(HexColor("#555555"))
    c.drawString(30 * mm, h - 35 * mm, "Project: Control Cabinet QC  |  Drawing: GA-001  |  Rev: A")
    c.drawString(30 * mm, h - 42 * mm, "Scale: 1:5  |  Date: 2026-10-12  |  Drawn: QC Inspector")

    cabinet_x = 60 * mm
    cabinet_y = 80 * mm
    cab_w = 280 * mm
    cab_h = 300 * mm

    c.setStrokeColor(HexColor("#333333"))
    c.setLineWidth(2)
    c.setFillColor(HexColor("#e8e8e8"))
    c.rect(cabinet_x, cabinet_y, cab_w, cab_h, fill=1, stroke=1)

    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(HexColor("#333333"))

    components = [
        ("PS11", "Power Supply\n24VDC / 10A", cabinet_x + 15 * mm, cabinet_y + cab_h - 60 * mm, 50 * mm, 45 * mm, "#4A90D9"),
        ("PS12", "Power Supply\n24VDC / 5A", cabinet_x + 75 * mm, cabinet_y + cab_h - 60 * mm, 50 * mm, 45 * mm, "#4A90D9"),
        ("D12", "Drive Unit\nVFD 1.5kW", cabinet_x + 140 * mm, cabinet_y + cab_h - 60 * mm, 55 * mm, 45 * mm, "#D94A4A"),
        ("CB41", "MCB\n16A C-curve", cabinet_x + 15 * mm, cabinet_y + cab_h - 115 * mm, 25 * mm, 40 * mm, "#5CB85C"),
        ("CB42", "MCB\n10A C-curve", cabinet_x + 45 * mm, cabinet_y + cab_h - 115 * mm, 25 * mm, 40 * mm, "#5CB85C"),
        ("CB43", "MCB\n6A C-curve", cabinet_x + 75 * mm, cabinet_y + cab_h - 115 * mm, 25 * mm, 40 * mm, "#5CB85C"),
        ("CB44", "MCB\n6A C-curve", cabinet_x + 105 * mm, cabinet_y + cab_h - 115 * mm, 25 * mm, 40 * mm, "#5CB85C"),
        ("K15", "Contactor\n9A AC3", cabinet_x + 145 * mm, cabinet_y + cab_h - 115 * mm, 40 * mm, 40 * mm, "#F0AD4E"),
        ("PLC-01", "PLC Module\nS7-1200", cabinet_x + 200 * mm, cabinet_y + cab_h - 115 * mm, 55 * mm, 40 * mm, "#337AB7"),
        ("X3", "Terminal Block\n12-way", cabinet_x + 15 * mm, cabinet_y + 25 * mm, 120 * mm, 25 * mm, "#777777"),
        ("KS2", "Safety Relay\nPNOZ", cabinet_x + 200 * mm, cabinet_y + cab_h - 170 * mm, 40 * mm, 35 * mm, "#D9534F"),
        ("1CBP03", "Baseplate", cabinet_x + cab_w / 2 - 30 * mm, cabinet_y + cab_h - 175 * mm, 60 * mm, 15 * mm, "#AAAAAA"),
    ]

    for label, desc, x, y, bw, bh, color in components:
        c.setFillColor(HexColor(color))
        c.setStrokeColor(HexColor("#333333"))
        c.setLineWidth(1)
        c.rect(x, y, bw, bh, fill=1, stroke=1)

        c.setFillColor(HexColor("#FFFFFF"))
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(x + bw / 2, y + bh - 12, label)
        c.setFont("Helvetica", 7)
        for i, line in enumerate(desc.split("\n")):
            c.drawCentredString(x + bw / 2, y + bh - 22 - i * 9, line)

    c.setStrokeColor(HexColor("#CC0000"))
    c.setLineWidth(1.5)
    c.setDash(3, 2)
    wire_routes = [
        (cabinet_x + 40 * mm, cabinet_y + cab_h - 60 * mm, cabinet_x + 27 * mm, cabinet_y + cab_h - 75 * mm),
        (cabinet_x + 100 * mm, cabinet_y + cab_h - 60 * mm, cabinet_x + 57 * mm, cabinet_y + cab_h - 75 * mm),
        (cabinet_x + 165 * mm, cabinet_y + cab_h - 115 * mm, cabinet_x + 165 * mm, cabinet_y + cab_h - 170 * mm),
    ]
    for x1, y1, x2, y2 in wire_routes:
        c.line(x1, y1, x2, y2)
    c.setDash()

    c.setFont("Helvetica", 7)
    c.setFillColor(HexColor("#999999"))
    c.drawString(cabinet_x, cabinet_y - 10 * mm,
                 "NOTE: All dimensions in mm. Wire colors per IEC 60445. Earth: Green/Yellow. Neutral: Blue. L+: Red/Brown.")

    c.setStrokeColor(HexColor("#333333"))
    c.setLineWidth(1.5)
    title_h = 20 * mm
    c.rect(30 * mm, 15 * mm, w - 60 * mm, title_h, stroke=1)
    c.setFont("Helvetica-Bold", 9)
    c.setFillColor(HexColor("#1a1a1a"))
    c.drawString(35 * mm, 28 * mm, "GA-001  |  CONTROL CABINET LAYOUT  |  CabinetSight QC Test")
    c.setFont("Helvetica", 8)
    c.drawString(35 * mm, 20 * mm, "Emerson Problem Statement #03  |  Project Eagle Eye  |  Sheet 1 of 1")

    c.save()
    print(f"  GA drawing: {out}")
    return out


if __name__ == "__main__":
    print("Generating CabinetSight test input files...")
    print(f"Output: {ASSETS}\n")

    create_label_list()
    create_wire_list()
    create_ga_drawing()

    photo = ROOT / "assets" / "cabinet-demo.png"
    print(f"\n  Photo: use {photo}")
    print(f"\nAll test files ready in: {ASSETS}")
    print("Upload these 4 files in the Workspace to run a real CV inspection.")
